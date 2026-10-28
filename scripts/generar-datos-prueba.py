#!/usr/bin/env python3
"""Genera libros Excel sintéticos reproducibles para demostraciones de ventas."""

from __future__ import annotations

import random
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Data" / "Demo_Ventas_2024_2025.xlsx"
HEADERS = [
    "Fecha", "Año", "Mes", "Punto", "Nombre_Punto", "Canal", "Región",
    "Producto_ID", "Producto", "Categoría", "Unidades", "Precio_Unitario_R$",
    "Ventas_R$", "Stock_Disponible_Unidades", "Lead_Time_Días", "Promoción",
    "Devoluciones_Unidades", "Ventas_Netas_R$", "Promedio_Diario_Aprox_Unid",
    "Demanda_30d_Estimada", "Stock_Seguridad_20pct", "Pedido_Sugerido_Base",
]
SEED = 20261028


def catalogues():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    products = [dict(zip((c.value for c in workbook["Catalogo_Productos"][1]), row))
                for row in workbook["Catalogo_Productos"].iter_rows(min_row=2, values_only=True)]
    points = [dict(zip((c.value for c in workbook["Puntos_Express"][1]), row))
              for row in workbook["Puntos_Express"].iter_rows(min_row=2, values_only=True)]
    workbook.close()
    return products, points


def generate(year: int, month_count: int, products: list[dict], points: list[dict], rng: random.Random):
    rows = []
    product_factor = {p["Producto_ID"]: rng.uniform(0.72, 1.32) for p in products}
    point_factor = {p["Punto"]: rng.uniform(0.72, 1.35) for p in points}
    for month in range(1, month_count + 1):
        for product in products:
            for point in points:
                pid, point_id = product["Producto_ID"], point["Punto"]
                # Base de demanda por producto y punto, con variación mensual acotada.
                base = 38 * product_factor[pid] * point_factor[point_id]
                season = 2.2 if month == 12 else 1.0 + 0.10 * __import__("math").sin((month - 1) * 2 * __import__("math").pi / 12)
                if year == 2026 and month in (6, 7) and pid in ("P001", "P004", "P011"):
                    season *= 1.8
                units = max(1, round(base * season * rng.uniform(0.86, 1.14)))

                promo = year == 2026 and rng.random() < 0.12
                if promo:
                    units = round(units * 1.4)
                price = round(product["Precio_Referencia_R$"] * rng.uniform(0.95, 1.05), 2)
                revenue = round(units * price, 2)
                returns = rng.randint(1, 3) if rng.random() < 0.10 else 0
                net_sales = round(revenue - returns * price, 2)
                demand = units
                safety = round(demand * 0.2)
                stock = max(0, round(demand * rng.uniform(0.85, 1.9)))

                # Quiebre demostrable al final de cada periodo para P001/P004/P007.
                if month >= month_count - 1 and pid in ("P001", "P004", "P007"):
                    stock = max(0, round(demand * (0.05 if month == month_count else 0.20)))
                # Sustitución: fresco agotado en el punto central, sustituto con inventario amplio.
                if year == 2026 and month == 10 and point_id == "SP-EXP-01":
                    if pid == "P001":
                        stock = 0
                    elif pid == "P002":
                        stock = demand * 4

                lead = 1 + (points.index(point) % 4)
                suggested = max(0, demand + safety - stock)
                rows.append([
                    datetime(year, month, 1), year, month, point_id, point["Nombre_Punto"],
                    point["Canal"], point["Región"], pid, product["Producto"], product["Categoría"],
                    units, price, revenue, stock, lead, "Sí" if promo else "No", returns,
                    net_sales, round(units / 30, 3), demand, safety, suggested,
                ])
    return rows


def write_book(path: Path, sheet_name: str, rows: list[list]):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for col in range(1, len(HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = min(
            34, max(13, max(len(str(sheet.cell(r, col).value or "")) for r in range(1, min(sheet.max_row, 100) + 1)) + 2)
        )
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "yyyy-mm-dd"
        for col in (12, 13, 18):
            sheet.cell(row, col).number_format = '0.00'
    workbook.save(path)


def main():
    products, points = catalogues()
    for year, months, name, sheet in (
        (2023, 12, "Test_Ventas_2023.xlsx", "Ventas_2023"),
        (2026, 10, "Test_Ventas_2026.xlsx", "Ventas_2026"),
    ):
        rows = generate(year, months, products, points, random.Random(SEED + year))
        output = ROOT / "Data" / name
        write_book(output, sheet, rows)
        print(f"{output.relative_to(ROOT)}: {len(rows)} filas")


if __name__ == "__main__":
    main()
